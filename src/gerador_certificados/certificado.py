import re
from io import BytesIO
from pathlib import Path
from pptx import Presentation
from pptx.shapes.group import GroupShape

TOKEN = re.compile(r"\{([A-Za-zÀ-ÿ0-9_ -]+)\}")

def _replace_in_paragraph(paragraph, replacements: dict[str, str]) -> None:
    original = "".join(run.text for run in paragraph.runs)
    if not original:
        return
    updated = TOKEN.sub(lambda match: replacements.get(match.group(1).strip().lower(), match.group(0)), original)
    if updated != original and paragraph.runs:
        paragraph.runs[0].text = updated
        for run in paragraph.runs[1:]:
            run.text = ""

def _replace_in_shape(shape, replacements: dict[str, str]) -> None:
    if isinstance(shape, GroupShape):
        for child in shape.shapes:
            _replace_in_shape(child, replacements)
        return
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for paragraph in cell.text_frame.paragraphs:
                    _replace_in_paragraph(paragraph, replacements)
        return
    if not shape.has_text_frame:
        return
    for paragraph in shape.text_frame.paragraphs:
        _replace_in_paragraph(paragraph, replacements)


def _tokens_da_forma(shape) -> set[str]:
    tokens: set[str] = set()
    if isinstance(shape, GroupShape):
        for child in shape.shapes:
            tokens.update(_tokens_da_forma(child))
        return tokens
    if shape.has_table:
        textos = [cell.text for row in shape.table.rows for cell in row.cells]
    elif shape.has_text_frame:
        textos = [shape.text]
    else:
        textos = []
    for texto in textos:
        tokens.update(token.strip().lower() for token in TOKEN.findall(texto))
    return tokens


def campos_modelo(modelo: Path) -> set[str]:
    """Retorna os marcadores {campo} encontrados nos slides e tabelas."""
    presentation = Presentation(modelo)
    return {
        campo
        for slide in presentation.slides
        for shape in slide.shapes
        for campo in _tokens_da_forma(shape)
    }

def preencher_modelo(
    modelo: Path,
    destino: Path,
    dados: dict[str, str],
    imagens: list[dict[str, float | str | bytes]] | None = None,
) -> None:
    """Preenche tokens {coluna} em todos os slides e salva uma cópia do modelo."""
    presentation = Presentation(modelo)
    replacements = {key.strip().lower(): value for key, value in dados.items()}
    for slide in presentation.slides:
        for shape in slide.shapes:
            _replace_in_shape(shape, replacements)
        for imagem in imagens or []:
            largura = int(presentation.slide_width * float(imagem.get("largura", 18)) / 100)
            esquerda = int(presentation.slide_width * float(imagem.get("x", 5)) / 100)
            topo = int(presentation.slide_height * float(imagem.get("y", 5)) / 100)
            slide.shapes.add_picture(BytesIO(imagem["dados"]), esquerda, topo, width=largura)
    presentation.save(destino)
