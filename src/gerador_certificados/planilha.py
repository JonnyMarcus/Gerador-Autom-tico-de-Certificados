from pathlib import Path
from openpyxl import load_workbook

def ler_participantes(caminho: Path, aba: str | None = None) -> list[dict[str, str]]:
    """Lê a primeira linha como cabeçalho e retorna registros normalizados."""
    if not caminho.is_file():
        raise FileNotFoundError(f"Planilha não encontrada: {caminho}")
    workbook = load_workbook(caminho, read_only=True, data_only=True)
    try:
        if aba and aba not in workbook.sheetnames:
            raise ValueError(f"Aba '{aba}' não encontrada. Abas disponíveis: {', '.join(workbook.sheetnames)}")
        sheet = workbook[aba] if aba else workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            raise ValueError("A planilha está vazia; inclua os nomes das colunas na primeira linha.")
        normalized = [str(value).strip().lower() if value is not None else "" for value in headers]
        if any(not header for header in normalized) or len(set(normalized)) != len(normalized):
            raise ValueError("Os cabeçalhos devem estar preenchidos e não podem se repetir.")
        participants = []
        for values in rows:
            if not any(value is not None and str(value).strip() for value in values):
                continue
            participants.append({header: "" if index >= len(values) or values[index] is None else str(values[index]).strip() for index, header in enumerate(normalized)})
        return participants
    finally:
        workbook.close()
