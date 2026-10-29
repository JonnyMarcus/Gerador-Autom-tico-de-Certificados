from pathlib import Path
from io import BytesIO
import subprocess
import sys
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from streamlit.testing.v1 import AppTest

from gerador_certificados.certificado import preencher_modelo
from gerador_certificados.config import ROOT_DIR
from gerador_certificados.exportador import nome_seguro
from gerador_certificados.pdf_certificado import (
    combinar_certificados_pdf,
    gerar_certificados_pdf,
)
from gerador_certificados.planilha import ler_participantes
from gerador_certificados.validacao import validar_participantes


class PlanilhaTests(unittest.TestCase):
    def test_ler_normaliza_cabecalhos_e_ignora_linhas_vazias(self):
        with TemporaryDirectory() as diretorio:
            arquivo = Path(diretorio) / "participantes.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append([" Nome ", "Curso"])
            sheet.append(["Ana Souza", "Primeiros socorros"])
            sheet.append([None, None])
            workbook.save(arquivo)

            self.assertEqual(
                ler_participantes(arquivo),
                [{"nome": "Ana Souza", "curso": "Primeiros socorros"}],
            )

    def test_rejeita_cabecalho_vazio_ou_duplicado(self):
        for headers in (["nome", None], ["nome", " NOME "]):
            with self.subTest(headers=headers), TemporaryDirectory() as diretorio:
                arquivo = Path(diretorio) / "participantes.xlsx"
                workbook = Workbook()
                workbook.active.append(headers)
                workbook.save(arquivo)
                with self.assertRaisesRegex(ValueError, "cabeçalhos"):
                    ler_participantes(arquivo)

    def test_rejeita_aba_inexistente(self):
        with TemporaryDirectory() as diretorio:
            arquivo = Path(diretorio) / "participantes.xlsx"
            Workbook().save(arquivo)
            with self.assertRaisesRegex(ValueError, "Aba 'Turma'"):
                ler_participantes(arquivo, "Turma")

    def test_validacao_exige_nome_e_rejeita_nome_vazio(self):
        with self.assertRaisesRegex(ValueError, "Nenhum participante"):
            validar_participantes([], "nome", [])
        with self.assertRaisesRegex(ValueError, "ausentes"):
            validar_participantes([{"curso": "A"}], "nome", [])
        with self.assertRaisesRegex(ValueError, "curso"):
            validar_participantes([{"nome": "Ana"}], "nome", ["curso"])
        with self.assertRaisesRegex(ValueError, "Nome vazio"):
            validar_participantes([{"nome": "Ana"}, {"nome": ""}], "nome", [])


class PdfTests(unittest.TestCase):
    @staticmethod
    def criar_modelo(caminho: Path, duas_paginas: bool = False) -> None:
        documento = canvas.Canvas(str(caminho), pagesize=(800, 500))
        documento.drawCentredString(400, 450, "CERTIFICADO DE CONCLUSÃO")
        documento.showPage()
        if duas_paginas:
            documento.drawCentredString(400, 450, "VERSO DO CERTIFICADO")
            documento.showPage()
        documento.save()

    def test_gera_um_pdf_por_participante_com_nome_e_fundo(self):
        with TemporaryDirectory() as diretorio:
            pasta = Path(diretorio)
            modelo = pasta / "modelo.pdf"
            self.criar_modelo(modelo)

            arquivos = gerar_certificados_pdf(
                modelo,
                pasta / "saida",
                [{"nome": "Ana Souza"}, {"nome": "João da Silva"}],
            )

            self.assertEqual(len(arquivos), 2)
            self.assertTrue(arquivos[0].exists())
            textos = [PdfReader(arquivo).pages[0].extract_text() for arquivo in arquivos]
            self.assertIn("CERTIFICADO DE CONCLUSÃO", textos[0])
            self.assertIn("Ana Souza", textos[0])
            self.assertIn("João da Silva", textos[1])
            self.assertNotIn("Ana Souza", textos[1])
            progresso = []
            gerar_certificados_pdf(
                modelo,
                pasta / "com_progresso",
                [{"nome": "Ana"}, {"nome": "João"}],
                ao_gerar=lambda atual, total: progresso.append((atual, total)),
            )
            self.assertEqual(progresso, [(1, 2), (2, 2)])

            pdf_unico = PdfReader(BytesIO(combinar_certificados_pdf(arquivos)))
            self.assertEqual(len(pdf_unico.pages), 2)
            self.assertIn("João da Silva", pdf_unico.pages[1].extract_text())

    def test_aceita_ajustes_visuais_e_valida_posicao(self):
        with TemporaryDirectory() as diretorio:
            pasta = Path(diretorio)
            modelo = pasta / "modelo.pdf"
            self.criar_modelo(modelo)
            arquivo = gerar_certificados_pdf(
                modelo,
                pasta / "saida",
                [{"nome": "Rita"}],
                posicao_x=25,
                posicao_y=30,
                fonte="Times-Bold",
                tamanho_fonte=24,
                cor="#C00000",
            )[0]
            self.assertIn("Rita", PdfReader(arquivo).pages[0].extract_text())
            with self.assertRaisesRegex(ValueError, "entre 0 e 100"):
                gerar_certificados_pdf(modelo, pasta / "invalido", [{"nome": "Rita"}], posicao_x=101)

    def test_nao_combina_lista_vazia(self):
        with self.assertRaisesRegex(ValueError, "Não há certificados"):
            combinar_certificados_pdf([])

    def test_preserva_paginas_adicionais_sem_escrever_nome_nelas(self):
        with TemporaryDirectory() as diretorio:
            pasta = Path(diretorio)
            modelo = pasta / "modelo.pdf"
            self.criar_modelo(modelo, duas_paginas=True)

            arquivo = gerar_certificados_pdf(modelo, pasta / "saida", [{"nome": "Marina"}])[0]
            paginas = PdfReader(arquivo).pages

            self.assertEqual(len(paginas), 2)
            self.assertIn("Marina", paginas[0].extract_text())
            self.assertIn("VERSO DO CERTIFICADO", paginas[1].extract_text())
            self.assertNotIn("Marina", paginas[1].extract_text())


class PowerPointTests(unittest.TestCase):
    def test_substitui_marcador_mesmo_quando_dividido_em_runs(self):
        with TemporaryDirectory() as diretorio:
            pasta = Path(diretorio)
            modelo = pasta / "modelo.pptx"
            saida = pasta / "saida.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
            paragraph = textbox.text_frame.paragraphs[0]
            paragraph.add_run().text = "Certificado de {no"
            paragraph.add_run().text = "me}; {curso}"
            presentation.save(modelo)

            preencher_modelo(modelo, saida, {"nome": "Ana", "curso": "Matemática"})
            texto = Presentation(saida).slides[0].shapes[0].text

            self.assertEqual(texto, "Certificado de Ana; Matemática")

    def test_substitui_marcador_dentro_de_tabela(self):
        with TemporaryDirectory() as diretorio:
            pasta = Path(diretorio)
            modelo = pasta / "modelo.pptx"
            saida = pasta / "saida.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            tabela = slide.shapes.add_table(1, 1, Inches(1), Inches(1), Inches(5), Inches(1)).table
            tabela.cell(0, 0).text = "{nome}"
            presentation.save(modelo)

            preencher_modelo(modelo, saida, {"nome": "Marina"})

            resultado = Presentation(saida)
            self.assertEqual(resultado.slides[0].shapes[0].table.cell(0, 0).text, "Marina")


class CliTests(unittest.TestCase):
    def test_cli_gera_arquivo_a_partir_de_excel_e_powerpoint(self):
        with TemporaryDirectory() as diretorio:
            pasta = Path(diretorio)
            planilha = pasta / "participantes.xlsx"
            modelo = pasta / "modelo.pptx"
            saida = pasta / "saida"

            workbook = Workbook()
            workbook.active.append(["nome", "curso"])
            workbook.active.append(["Ana Souza", "Matemática"])
            workbook.save(planilha)

            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = "{nome} · {curso}"
            presentation.save(modelo)

            resultado = subprocess.run(
                [
                    sys.executable,
                    "-m",
                    "gerador_certificados",
                    "--planilha",
                    str(planilha),
                    "--modelo",
                    str(modelo),
                    "--saida",
                    str(saida),
                ],
                cwd=ROOT_DIR,
                capture_output=True,
                text=True,
                check=False,
            )

            self.assertEqual(resultado.returncode, 0, resultado.stderr)
            arquivos = list(saida.glob("*.pptx"))
            self.assertEqual(len(arquivos), 1)
            self.assertEqual(Presentation(arquivos[0]).slides[0].shapes[0].text, "Ana Souza · Matemática")


class WebAppTests(unittest.TestCase):
    def test_interface_inicia_sem_erros(self):
        app = AppTest.from_file(str(ROOT_DIR / "app.py"), default_timeout=30).run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.get("file_uploader")), 2)
        self.assertEqual(len(app.get("text_input")), 0)

    def test_estilo_pronto_e_ajuste_manual_sincronizam(self):
        app = AppTest.from_file(str(ROOT_DIR / "app.py"), default_timeout=30).run()
        app.radio[0].set_value("Clássico").run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.selectbox[0].value, "Clássica · forte")
        self.assertEqual(app.slider[2].value, 34)

        app.slider[0].set_value(42).run()

        self.assertEqual(len(app.exception), 0)
        self.assertEqual(app.radio[0].value, "Personalizado")

    def test_upload_e_geracao_disponibilizam_as_duas_opcoes_de_download(self):
        workbook_buffer = BytesIO()
        workbook = Workbook()
        workbook.active.append(["nome"])
        workbook.active.append(["Ana Souza"])
        workbook.save(workbook_buffer)

        pdf_buffer = BytesIO()
        documento = canvas.Canvas(pdf_buffer, pagesize=(800, 500))
        documento.drawCentredString(400, 450, "CERTIFICADO")
        documento.showPage()
        documento.save()

        app = AppTest.from_file(str(ROOT_DIR / "app.py"), default_timeout=30).run()
        uploaders = app.get("file_uploader")
        uploaders[0].set_value(("participantes.xlsx", workbook_buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        uploaders[1].set_value(("modelo.pdf", pdf_buffer.getvalue(), "application/pdf"))
        app.run()
        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.get("iframe")), 2)
        previa_inicial = app.session_state["pdf_previa"]
        app.slider[0].set_value(42).run()
        self.assertNotEqual(app.session_state["pdf_previa"], previa_inicial)

        app.button[0].click().run()

        self.assertEqual(len(app.exception), 0)
        self.assertEqual(len(app.get("download_button")), 2)


class NomeSeguroTests(unittest.TestCase):
    def test_remove_acentos_e_caracteres_de_caminho(self):
        self.assertEqual(nome_seguro("João / Silva"), "Joao_Silva")
        self.assertEqual(nome_seguro("../"), "participante")


if __name__ == "__main__":
    unittest.main()
