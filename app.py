"""Interface web do gerador, preparada para Streamlit Community Cloud."""

from __future__ import annotations

import base64
from io import BytesIO
from hashlib import sha256
from pathlib import Path
import sys
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile
import os
import uuid
from datetime import datetime
from zoneinfo import ZoneInfo

import streamlit as st
import pymupdf as fitz
import qrcode
from PIL import Image
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import landscape, A4
from reportlab.pdfgen import canvas as pdf_canvas

# O Streamlit Cloud executa app.py na raiz; o pacote da aplicação fica em src/.
sys.path.insert(0, str(Path(__file__).resolve().parent / "src"))

from gerador_certificados.planilha import ler_participantes
from gerador_certificados.pdf_certificado import (
    combinar_certificados_pdf,
    gerar_certificados_pdf,
)
from gerador_certificados.validacao import validar_participantes
from gerador_certificados.posicao_component import editor_posicao
from gerador_certificados.exportador import nome_seguro
from gerador_certificados.certificado import preencher_modelo
from gerador_certificados.seguranca_qr import gerar_link_verificacao, verificar_link_verificacao
from gerador_certificados.certificado import campos_modelo


st.set_page_config(page_title="Certifica | Gerador de certificados", page_icon="🎓", layout="wide")

MAX_XLSX_MB = 5
MAX_PDF_MB = 10
MAX_IMAGEM_MB = 5
MAX_PARTICIPANTES = 500
MAX_SAIDA_ESTIMADA_BYTES = 100 * 1024 * 1024
FONTES_EXIBICAO = {
    "Moderna · regular": "Helvetica",
    "Moderna · negrito": "Helvetica-Bold",
    "Moderna · itálico": "Helvetica-Oblique",
    "Moderna · negrito itálico": "Helvetica-BoldOblique",
    "Clássica · regular": "Times-Roman",
    "Clássica · negrito": "Times-Bold",
    "Clássica · itálico": "Times-Italic",
    "Clássica · negrito itálico": "Times-BoldItalic",
    "Monoespaçada · regular": "Courier",
    "Monoespaçada · negrito": "Courier-Bold",
    "Monoespaçada · itálico": "Courier-Oblique",
    "Monoespaçada · negrito itálico": "Courier-BoldOblique",
    "Manuscrita · Great Vibes": "GreatVibes",
}
ESTILOS = {
    "Moderno": {"x": 50, "y": 51, "fonte": "Helvetica-Bold", "tamanho": 32, "cor": "#17233C"},
    "Clássico": {"x": 50, "y": 52, "fonte": "Times-Bold", "tamanho": 34, "cor": "#795B32"},
    "Institucional": {"x": 50, "y": 51, "fonte": "Helvetica-Bold", "tamanho": 28, "cor": "#5148E5"},
    "Acadêmico": {"x": 50, "y": 52, "fonte": "Times-Bold", "tamanho": 32, "cor": "#243D63"},
    "Corporativo": {"x": 50, "y": 50, "fonte": "Helvetica-Bold", "tamanho": 30, "cor": "#1C3557"},
    "Minimalista": {"x": 50, "y": 51, "fonte": "Helvetica", "tamanho": 28, "cor": "#30343B"},
}


def obter_configuracao_qr(chave: str, padrao: str = "") -> str:
    try:
        valor = st.secrets.get(chave, padrao)
        return str(valor) if valor else padrao
    except Exception:
        return os.environ.get(f"CERTIFICA_{chave.upper()}", padrao)


def criar_qr_png(dados: dict[str, object], segredo: str, url_publica: str) -> bytes:
    codigo = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_M, box_size=5, border=2)
    codigo.add_data(gerar_link_verificacao(dados, segredo, url_publica))
    codigo.make(fit=True)
    imagem = codigo.make_image(fill_color="#17233C", back_color="white")
    buffer = BytesIO()
    imagem.save(buffer, format="PNG")
    return buffer.getvalue()


def criar_dados_verificacao(
    participante: dict[str, str], formato: str, data_emissao: str, campos: list[str]
) -> dict[str, object]:
    return {
        "id": uuid.uuid4().hex[:16].upper(),
        "nome": participante.get("nome", ""),
        "emissao": data_emissao,
        "formato": formato,
        "campos": {campo: participante.get(campo, "") for campo in campos},
    }


def qr_participante_png(
    participante: dict[str, str], formato: str, emissao: str, campos: list[str]
) -> bytes:
    dados = criar_dados_verificacao(participante, formato, emissao, campos)
    return criar_qr_png(dados, SEGREDO_QR, URL_PUBLICA_QR)


def preparar_imagem(arquivo, x: int, y: int, largura: int) -> dict[str, float | str | bytes]:
    dados = arquivo.getvalue()
    with Image.open(BytesIO(dados)) as imagem:
        if imagem.width > 6000 or imagem.height > 6000 or imagem.width * imagem.height > 30_000_000:
            raise ValueError("A imagem deve ter no máximo 6000 px por lado e 30 megapixels.")
        imagem.verify()
    return {"dados": dados, "x": x, "y": y, "largura": largura}


def guardar_no_historico(formato: str, quantidade: int, arquivo_zip: bytes) -> None:
    limite_bytes = 60 * 1024 * 1024
    historico = st.session_state.get("historico_emissoes", [])
    payload = arquivo_zip if len(arquivo_zip) <= limite_bytes else None
    while historico and sum(len(item.get("zip") or b"") for item in historico) + len(payload or b"") > limite_bytes:
        historico.pop()
    historico.insert(
        0,
        {
            "id": uuid.uuid4().hex[:8],
            "data": datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M"),
            "formato": formato,
            "quantidade": quantidade,
            "zip": payload,
        },
    )
    st.session_state["historico_emissoes"] = historico[:3]


SEGREDO_QR = obter_configuracao_qr("qr_secret")
URL_PUBLICA_QR = obter_configuracao_qr(
    "public_url", "https://arpalq47szetr6bwvpxtzy.streamlit.app"
)
TOKEN_VERIFICACAO = st.query_params.get("verificar")
if TOKEN_VERIFICACAO:
    st.title("Verificação de certificado")
    dados_verificados = verificar_link_verificacao(TOKEN_VERIFICACAO, SEGREDO_QR) if len(SEGREDO_QR) >= 32 else None
    if dados_verificados:
        st.success("QR Code válido: os dados conferem com a assinatura digital do certificado.")
        st.markdown(f"**Certificado:** {dados_verificados.get('id', '—')}")
        st.markdown(f"**Participante:** {dados_verificados.get('nome', '—')}")
        st.markdown(f"**Emissão:** {dados_verificados.get('emissao', '—')}")
        st.markdown(f"**Formato:** {dados_verificados.get('formato', '—')}")
        if dados_verificados.get("campos"):
            st.json(dados_verificados["campos"])
    else:
        st.error("Este QR Code não é válido ou foi assinado com outra chave.")
    st.stop()


def criar_zip(arquivos: list[Path]) -> bytes:
    buffer = BytesIO()
    with ZipFile(buffer, "w", compression=ZIP_DEFLATED) as pacote:
        for arquivo in arquivos:
            pacote.write(arquivo, arcname=arquivo.name)
    return buffer.getvalue()


def gerar_previa_pdf(
    modelo_pdf: bytes,
    participante: dict[str, str],
    posicao_x: int,
    posicao_y: int,
    fonte: str,
    tamanho_fonte: int,
    cor: str,
    campos_extras: dict[str, dict[str, float | str]] | None = None,
    imagens: list[dict[str, float | str | bytes]] | None = None,
    qr_png: bytes | None = None,
    qr_posicao: dict[str, float] | None = None,
) -> bytes:
    with tempfile.TemporaryDirectory(prefix="certificados_previa_") as diretorio:
        pasta = Path(diretorio)
        modelo = pasta / "modelo.pdf"
        modelo.write_bytes(modelo_pdf)
        arquivo = gerar_certificados_pdf(
            modelo,
            pasta / "saida",
            [participante],
            posicao_x=posicao_x,
            posicao_y=posicao_y,
            fonte=fonte,
            tamanho_fonte=tamanho_fonte,
            cor=cor,
            campos_extras=campos_extras,
            imagens=imagens,
            qr_por_participante=(lambda _participante: qr_png) if qr_png else None,
            qr_posicao=qr_posicao,
        )[0]
        return arquivo.read_bytes()


def imagem_primeira_pagina(pdf: bytes) -> str:
    """Renderiza a primeira página para o editor visual de posição."""
    documento = fitz.open(stream=pdf, filetype="pdf")
    try:
        if documento.needs_pass:
            raise ValueError("O PDF está protegido por senha. Envie uma cópia desbloqueada.")
        if not documento.page_count:
            raise ValueError("O PDF não contém páginas.")
        pagina = documento[0]
        if pagina.rect.width < 100 or pagina.rect.height < 100:
            raise ValueError("A primeira página do PDF tem dimensões muito pequenas.")
        escala = min(1.35, 1400 / max(pagina.rect.width, pagina.rect.height))
        pixmap = pagina.get_pixmap(matrix=fitz.Matrix(escala, escala), alpha=False)
        return "data:image/png;base64," + base64.b64encode(pixmap.tobytes("png")).decode("ascii")
    finally:
        documento.close()


def diagnosticar_area_pdf(pdf: bytes, areas: dict[str, tuple[float, float]]) -> tuple[str, list[str]]:
    """Resume a página e sinaliza texto detectável perto das posições de impressão."""
    documento = fitz.open(stream=pdf, filetype="pdf")
    try:
        if documento.needs_pass:
            raise ValueError("O PDF está protegido por senha. Envie uma cópia desbloqueada.")
        if not documento.page_count:
            raise ValueError("O PDF não contém páginas.")
        pagina = documento[0]
        largura, altura = pagina.rect.width, pagina.rect.height
        resumo = f"{documento.page_count} página(s) · primeira página {largura:.0f} × {altura:.0f} pt"
        palavras = pagina.get_text("words")
        avisos = []
        for rotulo, (x, y) in areas.items():
            centro_y = altura * y / 100
            tolerancia_y = max(16, altura * 0.035)
            tem_texto = any(
                palavra[0] < largura * 0.92
                and palavra[2] > largura * 0.08
                and palavra[1] <= centro_y + tolerancia_y
                and palavra[3] >= centro_y - tolerancia_y
                for palavra in palavras
            )
            if tem_texto:
                avisos.append(
                    f"Há texto selecionável próximo à área de {rotulo} (vertical {y:.0f}%). "
                    "Confira a prévia e mova o campo se houver sobreposição."
                )
        return resumo, avisos
    finally:
        documento.close()


def criar_planilha_exemplo() -> bytes:
    arquivo = BytesIO()
    pasta = Workbook()
    planilha = pasta.active
    planilha.title = "Participantes"
    planilha.append(["nome", "curso", "carga_horaria", "data"])
    planilha.append(["Ana Souza", "Primeiros Socorros", "8 horas", "28/10/2026"])
    planilha.append(["João Lima", "Primeiros Socorros", "8 horas", "28/10/2026"])
    for coluna, largura in {"A": 28, "B": 32, "C": 20, "D": 18}.items():
        planilha.column_dimensions[coluna].width = largura
    pasta.save(arquivo)
    return arquivo.getvalue()


def criar_pdf_exemplo() -> bytes:
    arquivo = BytesIO()
    largura, altura = landscape(A4)
    pdf = pdf_canvas.Canvas(arquivo, pagesize=(largura, altura))
    pdf.setFillColor(HexColor("#FBFAF6"))
    pdf.rect(0, 0, largura, altura, stroke=0, fill=1)
    pdf.setStrokeColor(HexColor("#B89A5A"))
    pdf.setLineWidth(2)
    pdf.roundRect(24, 24, largura - 48, altura - 48, 14, stroke=1, fill=0)
    pdf.setLineWidth(.6)
    pdf.roundRect(34, 34, largura - 68, altura - 68, 10, stroke=1, fill=0)
    pdf.setFillColor(HexColor("#25314A"))
    pdf.setFont("Times-Bold", 30)
    pdf.drawCentredString(largura / 2, altura * .75, "CERTIFICADO DE CONCLUSÃO")
    pdf.setFillColor(HexColor("#66748B"))
    pdf.setFont("Helvetica", 13)
    pdf.drawCentredString(largura / 2, altura * .64, "Certificamos que")
    pdf.setStrokeColor(HexColor("#D7C79F"))
    pdf.line(largura * .18, altura * .48, largura * .82, altura * .48)
    pdf.setFillColor(HexColor("#8A7955"))
    pdf.setFont("Helvetica-Oblique", 10)
    pdf.drawCentredString(largura / 2, altura * .44, "A área central está livre para nome e dados personalizados")
    pdf.setFillColor(HexColor("#66748B"))
    pdf.setFont("Helvetica", 10)
    pdf.drawString(largura * .2, altura * .2, "Responsável")
    pdf.drawRightString(largura * .8, altura * .2, "Data")
    pdf.save()
    return arquivo.getvalue()


st.markdown(
    """
    <style>
      :root { --ink:#17233c; --muted:#66748b; --brand:#5148e5; --mint:#22b8a5; }
      [data-testid="stAppViewContainer"] { background:linear-gradient(180deg,#f3f6fc 0%,#fbfcff 340px,#fbfcff 100%); }
      [data-testid="stHeader"] { background:transparent; }
      .block-container { max-width:1120px; padding-top:2rem; padding-bottom:4rem; }
      .step-card { min-height:70px; display:flex; align-items:center; gap:.7rem; background:#fff; border:1px solid #e5eaf2; border-radius:14px; padding:.8rem 1rem; color:#66748b; font-size:.84rem; margin:.35rem 0 1.2rem; }
      .step-card strong { display:block; color:#17233c; font-size:.92rem; }
      .step-icon { width:30px; height:30px; flex:0 0 30px; display:grid; place-items:center; border-radius:50%; background:#edf0f6; color:#66748b; font-weight:800; }
      .step-card.complete { border-color:#b3eadc; background:#f3fcf9; }
      .step-card.complete .step-icon { background:#d8f5ec; color:#16866f; }
      .step-card.active { border-color:#b9b5ff; background:#f4f3ff; color:#5148e5; }
      .step-card.active .step-icon { background:#5148e5; color:#fff; }
      .step-card.pending { opacity:.78; }
      .section-title { color:#17233c; font-size:1.25rem; font-weight:750; margin:.7rem 0 .2rem; }
      .section-help { color:#718096; font-size:.9rem; margin-bottom:1rem; }
      [data-testid="stFileUploader"] section { background:#fff; border:1px dashed #bac6d8; border-radius:15px; }
      [data-testid="stFileUploader"] section:hover { border-color:#5148e5; }
      [data-testid="stExpander"] { background:#fff; border:1px solid #e6eaf1; border-radius:14px; }
      div.stButton > button, div.stDownloadButton > button { min-height:3rem; border-radius:12px; font-weight:700; }
      [data-testid="stDataFrame"] { border:1px solid #e5eaf2; border-radius:12px; overflow:hidden; }
      .privacy-note { color:#7b8798; font-size:.78rem; text-align:center; margin-top:2rem; }
      @media (max-width: 700px) { .block-container { padding:1.1rem 1rem 3rem; } }
    </style>
    """,
    unsafe_allow_html=True,
)

animar_entrada = not st.session_state.get("entrada_animada", False)
st.session_state["entrada_animada"] = True

st.iframe(
    """
    <!doctype html><html lang="pt-BR"><head><meta charset="utf-8">
    <script src="https://cdn.jsdelivr.net/npm/gsap@3.13.0/dist/gsap.min.js"></script>
    <style>
      *{box-sizing:border-box} html,body{margin:0;padding:0;background:transparent;font-family:Inter,ui-sans-serif,system-ui,-apple-system,"Segoe UI",sans-serif}
      .hero{height:326px;position:relative;overflow:hidden;border-radius:26px;padding:36px 44px;color:#fff;background:radial-gradient(ellipse at 83% 12%,rgba(123,115,255,.48),transparent 35%),linear-gradient(118deg,#17233c 0%,#253865 57%,#5148e5 100%);box-shadow:0 18px 42px rgba(38,53,94,.16)}
      .hero:before,.hero:after{content:"";position:absolute;border:1px solid rgba(255,255,255,.1);border-radius:50%;pointer-events:none}
      .hero:before{width:370px;height:370px;right:45px;top:-155px}.hero:after{width:500px;height:500px;right:-20px;top:-215px}
      .hero-copy{position:relative;z-index:2;width:58%;padding-top:7px}.eyebrow{display:inline-flex;align-items:center;gap:8px;padding:7px 11px;border:1px solid rgba(255,255,255,.2);border-radius:30px;background:rgba(255,255,255,.1);color:#e7eaff;font-size:11px;font-weight:750;letter-spacing:.12em;text-transform:uppercase}
      .eyebrow i{width:7px;height:7px;border-radius:50%;background:#50e0c2;box-shadow:0 0 12px #50e0c2}.title{margin:17px 0 12px;font-size:clamp(30px,4vw,43px);line-height:1.08;letter-spacing:-.045em;font-weight:760}.title-line{display:block;overflow:hidden}.title-line span{display:block}.copy{max-width:450px;color:#d8e2ff;font-size:14px;line-height:1.65;margin:0}.hero-foot{display:flex;align-items:center;gap:9px;margin-top:20px;color:#becbf0;font-size:11px}.hero-foot b{color:#fff;font-weight:650}.spark{color:#58e0c5;font-size:16px}
      .art{position:absolute;z-index:1;right:7.5%;top:47px;width:250px;height:230px;display:grid;place-items:center;perspective:800px}.orbit{position:absolute;inset:0;border:1px solid rgba(255,255,255,.16);border-radius:50%;transform:rotate(-20deg) scaleY(.54)}.orbit:after{content:"";position:absolute;left:18%;top:7%;width:8px;height:8px;border-radius:50%;background:#64e3ce;box-shadow:0 0 18px #64e3ce}
      .certificate{position:relative;width:208px;height:148px;padding:16px 17px;border:1px solid rgba(255,255,255,.8);border-radius:12px;background:linear-gradient(145deg,#fff,#edf2ff);box-shadow:0 22px 48px rgba(8,16,42,.32);transform:rotate(5deg)}.cert-top{display:flex;justify-content:space-between;align-items:center;color:#697799;font-size:7px;letter-spacing:.1em;text-transform:uppercase}.seal{width:21px;height:21px;border:1px solid #776ee7;border-radius:50%;display:grid;place-items:center;color:#5148e5;font-size:10px}.cert-label{text-align:center;color:#5148e5;font-size:7px;font-weight:800;letter-spacing:.16em;margin-top:13px}.cert-name{text-align:center;color:#182647;font:italic 19px Georgia,serif;margin-top:5px}.cert-rule{height:1px;width:75%;margin:6px auto;background:#cbd3e5}.cert-copy{text-align:center;color:#7d88a0;font-size:6px;line-height:1.5}.cert-bottom{display:flex;justify-content:space-between;margin-top:11px;color:#7d88a0;font-size:6px}.cert-bottom span{border-top:1px solid #c8d0e0;padding-top:3px;width:61px;text-align:center}
      .float-chip{position:absolute;padding:8px 11px;border:1px solid rgba(255,255,255,.18);border-radius:10px;background:rgba(255,255,255,.12);backdrop-filter:blur(8px);color:#fff;font-size:9px;font-weight:650;box-shadow:0 9px 22px rgba(8,16,42,.15)}.chip-one{right:0;top:32px}.chip-two{left:0;bottom:26px}.check{color:#55e0c2;margin-right:5px}
      @media(max-width:700px){.hero{height:326px;padding:26px 23px;border-radius:20px}.hero-copy{width:100%;padding-top:1px}.title{font-size:32px;max-width:360px}.copy{font-size:13px;max-width:350px}.art{width:150px;height:125px;right:7%;top:auto;bottom:-8px;opacity:.82}.certificate{transform:scale(.77) rotate(5deg)}.orbit{inset:0}.chip-one{right:0;top:9px}.chip-two{left:-5px;bottom:2px}}
      @media(prefers-reduced-motion:reduce){*{scroll-behavior:auto!important}}
    </style></head><body>
    <section class="hero">
      <div class="hero-copy">
        <div class="eyebrow"><i></i> Certifica · emissão simplificada</div>
        <h1 class="title"><span class="title-line"><span>Transforme sua lista</span></span><span class="title-line"><span>em certificados prontos.</span></span></h1>
        <p class="copy">Envie a lista de participantes e seu certificado em PDF. Cada nome vira um certificado pronto para compartilhar.</p>
        <div class="hero-foot"><span class="spark">✦</span><span><b>Simples de preparar.</b> Fácil de compartilhar.</span></div>
      </div>
      <div class="art"><div class="orbit"></div><div class="certificate"><div class="cert-top"><span>Certifica · 2026</span><span class="seal">✦</span></div><div class="cert-label">CERTIFICADO DE CONCLUSÃO</div><div class="cert-name">Seu participante</div><div class="cert-rule"></div><div class="cert-copy">concluiu com êxito o programa<br>de formação e desenvolvimento</div><div class="cert-bottom"><span>Coordenação</span><span>Data de emissão</span></div></div><div class="float-chip chip-one"><span class="check">✓</span>Personalizado</div><div class="float-chip chip-two">PDF pronto para compartilhar</div></div>
    </section>
    <script>
      const reduceMotion=window.matchMedia('(prefers-reduced-motion: reduce)').matches;
      const playIntro=__PLAY_INTRO__;
      if(window.gsap&&!reduceMotion&&playIntro){
        const intro=gsap.timeline({defaults:{ease:'power3.out'}});
        intro.from('.hero',{autoAlpha:0,y:20,duration:.65})
          .from('.eyebrow',{autoAlpha:0,y:12,duration:.38},'-=.28')
          .from('.title-line span',{yPercent:115,duration:.72,stagger:.12},'-=.18')
          .from('.copy,.hero-foot',{autoAlpha:0,y:12,duration:.45,stagger:.1},'-=.35')
          .from('.certificate',{autoAlpha:0,y:22,rotateY:-18,rotateZ:-3,duration:.8},'-=.72')
          .from('.float-chip',{autoAlpha:0,y:10,scale:.92,duration:.35,stagger:.1},'-=.45');
        gsap.to('.certificate',{y:-7,rotateZ:3,duration:2.8,ease:'sine.inOut',repeat:-1,yoyo:true,delay:1.4});
        gsap.to('.orbit',{rotate:0,duration:22,ease:'none',repeat:-1});
        gsap.to('.eyebrow i',{scale:1.35,opacity:.65,duration:.9,repeat:-1,yoyo:true,ease:'sine.inOut'});
      }
    </script></body></html>
    """.replace("__PLAY_INTRO__", str(animar_entrada).lower()),
    width="stretch",
    height=350,
    alt="Apresentação animada do gerador de certificados",
)

with st.expander("Antes de começar · confira o formato dos arquivos", expanded=False):
    st.markdown(
        """
        **Planilha:** arquivo `.xlsx`, com os cabeçalhos na primeira linha e uma coluna chamada `nome`.

        **Certificado:** arquivo `.pdf`. O nome será aplicado ao centro da primeira página.

        Deixe uma área livre no centro do PDF para o nome do participante.
        """
    )

modelo_coluna, certificado_coluna = st.columns(2)
with modelo_coluna:
    st.download_button(
        "Baixar modelo de planilha (.xlsx)",
        data=criar_planilha_exemplo(),
        file_name="modelo_participantes.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
    )
with certificado_coluna:
    st.download_button(
        "Baixar certificado de exemplo (.pdf)",
        data=criar_pdf_exemplo(),
        file_name="certificado_exemplo.pdf",
        mime="application/pdf",
        width="stretch",
    )

formato_saida = st.radio(
    "Formato do modelo e dos certificados",
    options=["PDF", "PowerPoint"],
    horizontal=True,
    key="formato_saida",
    help="O PDF recebe os ajustes visuais nesta tela. No PowerPoint, os marcadores {nome}, {curso} etc. são substituídos e a formatação vem do modelo.",
)
tipo_modelo = ["pdf"] if formato_saida == "PDF" else ["pptx"]
rotulo_modelo = "Certificado base em PDF" if formato_saida == "PDF" else "Modelo PowerPoint (.pptx)"
arquivo_modelo = None

uploaded_count = 0

left_step, right_step = st.columns(2)
with left_step:
    st.markdown('<div class="section-title">1. Lista de participantes</div><div class="section-help">Selecione a planilha Excel (.xlsx)</div>', unsafe_allow_html=True)
    arquivo_planilha = st.file_uploader("Escolher planilha", type=["xlsx"], label_visibility="collapsed", key="planilha_upload", max_upload_size=MAX_XLSX_MB)
    if arquivo_planilha:
        uploaded_count += 1
        st.caption(f"✓ {arquivo_planilha.name} · {arquivo_planilha.size / (1024 * 1024):.2f} MB")
with right_step:
    st.markdown(f'<div class="section-title">2. {rotulo_modelo}</div><div class="section-help">Envie o arquivo do modelo</div>', unsafe_allow_html=True)
    arquivo_modelo = st.file_uploader(
        "Escolher modelo de certificado",
        type=tipo_modelo,
        label_visibility="collapsed",
        key="pdf_upload" if formato_saida == "PDF" else "pptx_upload",
        max_upload_size=MAX_PDF_MB,
    )
    if arquivo_modelo:
        uploaded_count += 1
        st.caption(f"✓ {arquivo_modelo.name} · {arquivo_modelo.size / (1024 * 1024):.2f} MB")

arquivo_pdf = arquivo_modelo if formato_saida == "PDF" else None
arquivo_pptx = arquivo_modelo if formato_saida == "PowerPoint" else None

steps = st.columns(3)
status_steps = [
    ("1", "Lista pronta" if arquivo_planilha else "Envie a lista", bool(arquivo_planilha)),
    ("2", f"{formato_saida} pronto" if arquivo_modelo else "Envie o modelo", bool(arquivo_modelo)),
    ("3", "Pronto para gerar" if uploaded_count == 2 else "Gere e baixe", False),
]
proximo_passo = 0 if not arquivo_planilha else (1 if not arquivo_modelo else 2)
for indice, (numero, texto, concluido) in enumerate(status_steps):
    classe = "complete" if concluido else ("active" if indice == proximo_passo else "pending")
    icone = "✓" if concluido else numero
    steps[indice].markdown(
        f'<div class="step-card {classe}"><span class="step-icon">{icone}</span><span><strong>Etapa {numero}</strong>{texto}</span></div>',
        unsafe_allow_html=True,
    )

coluna_nome = "nome"
amostra = []
erro_dados = None
erro_modelo_pptx = None
marcadores_modelo = set()
if arquivo_planilha:
    try:
        with tempfile.TemporaryDirectory(prefix="certificados_lista_") as diretorio:
            caminho_planilha = Path(diretorio) / "participantes.xlsx"
            caminho_planilha.write_bytes(arquivo_planilha.getvalue())
            amostra = ler_participantes(caminho_planilha)
        if amostra:
            st.markdown(
                f'<div class="section-title">Prévia da planilha '
                f'<span style="color:#5148e5;font-size:.9rem">· {len(amostra)} participantes</span></div>',
                unsafe_allow_html=True,
            )
            st.dataframe(amostra[:10], width="stretch", hide_index=True)
            contagem_nomes = {}
            for indice, participante in enumerate(amostra, start=2):
                chave_nome = participante.get(coluna_nome, "").strip().casefold()
                if chave_nome:
                    contagem_nomes.setdefault(chave_nome, []).append(indice)
            repetidos = [linhas for linhas in contagem_nomes.values() if len(linhas) > 1]
            if repetidos:
                linhas_repetidas = ", ".join(str(linha) for linhas in repetidos for linha in linhas)
                st.warning(
                    f"Há nomes repetidos nas linhas {linhas_repetidas}. "
                    "A geração continuará; confira se as repetições são intencionais."
                )
    except (ValueError, OSError, KeyError) as error:
        erro_dados = str(error)
        st.error(f"Não foi possível ler a planilha: {error}")

campos_selecionados = []
if formato_saida == "PDF":
    campos_disponiveis = [campo for campo in amostra[0] if campo != coluna_nome] if amostra else []
    st.session_state["campos_extras_selecionados"] = [
        campo for campo in st.session_state.get("campos_extras_selecionados", []) if campo in campos_disponiveis
    ]
    campos_selecionados = st.multiselect(
        "Dados adicionais para imprimir no certificado",
        options=campos_disponiveis,
        help="Escolha colunas como curso, carga_horaria ou data. Depois ajuste a posição de cada uma.",
        key="campos_extras_selecionados",
    )
    if amostra and campos_selecionados:
        try:
            validar_participantes(amostra, coluna_nome, campos_selecionados)
        except ValueError as error:
            erro_dados = str(error)
            st.error(f"Confira os dados selecionados: {error}")
elif arquivo_pptx and amostra:
    try:
        with tempfile.TemporaryDirectory(prefix="certificados_modelo_pptx_") as diretorio:
            caminho_modelo_pptx = Path(diretorio) / "modelo.pptx"
            caminho_modelo_pptx.write_bytes(arquivo_pptx.getvalue())
            marcadores_modelo = campos_modelo(caminho_modelo_pptx)
        if "nome" not in marcadores_modelo:
            erro_modelo_pptx = "O modelo PowerPoint não contém o marcador obrigatório {nome}."
            st.error(erro_modelo_pptx)
        faltantes_modelo = marcadores_modelo - set(amostra[0])
        if faltantes_modelo:
            erro_modelo_pptx = f"Faltam colunas na planilha para estes marcadores: {', '.join(sorted(faltantes_modelo))}."
            st.error(erro_modelo_pptx)
        elif not erro_modelo_pptx:
            st.info("O PowerPoint preservará o visual do modelo e preencherá seus marcadores {coluna}.")
    except Exception as error:
        erro_modelo_pptx = f"Não foi possível ler o modelo PowerPoint: {error}"
        st.error(erro_modelo_pptx)

def marcar_personalizado() -> None:
    st.session_state["estilo_personalizacao"] = "Personalizado"


posicao_x = 50
posicao_y = 51
fonte = "Helvetica-Bold"
tamanho_fonte = 32
cor_texto = "#17233C"
campos_extras = {}
if formato_saida == "PDF":
    st.session_state.setdefault("estilo_personalizacao", "Personalizado")
    st.session_state.setdefault("personalizacao_x", 50)
    st.session_state.setdefault("personalizacao_y", 51)
    st.session_state.setdefault("personalizacao_fonte", "Helvetica-Bold")
    st.session_state.setdefault("personalizacao_tamanho", 32)
    st.session_state.setdefault("personalizacao_cor", "#17233C")
    st.session_state.setdefault("estilo_aplicado", "Personalizado")

    with st.expander("Personalize o nome", expanded=True):
        st.markdown("Comece por um estilo pronto ou ajuste cada detalhe. A prévia acompanha suas mudanças.")
        estilo = st.radio(
            "Estilo inicial",
            options=["Personalizado", *ESTILOS],
            horizontal=True,
            key="estilo_personalizacao",
        )
        if estilo != st.session_state["estilo_aplicado"]:
            if estilo in ESTILOS:
                preset = ESTILOS[estilo]
                st.session_state["personalizacao_x"] = preset["x"]
                st.session_state["personalizacao_y"] = preset["y"]
                st.session_state["personalizacao_fonte"] = preset["fonte"]
                st.session_state["rotulo_fonte"] = next(
                    rotulo for rotulo, fonte_pdf in FONTES_EXIBICAO.items() if fonte_pdf == preset["fonte"]
                )
                st.session_state["personalizacao_tamanho"] = preset["tamanho"]
                st.session_state["personalizacao_cor"] = preset["cor"]
            st.session_state["estilo_aplicado"] = estilo

        st.session_state.setdefault("rotulo_fonte", "Moderna · negrito")

        esquerda, direita = st.columns(2)
        with esquerda:
            posicao_x = st.slider(
                "Posição horizontal",
                min_value=10,
                max_value=90,
                key="personalizacao_x",
                help="Controla o centro do nome entre as margens.",
                on_change=marcar_personalizado,
            )
            posicao_y = st.slider(
                "Posição vertical",
                min_value=10,
                max_value=90,
                key="personalizacao_y",
                help="10% fica perto do topo; 90% fica perto da parte inferior.",
                on_change=marcar_personalizado,
            )
        with direita:
            st.markdown("**Estilo tipográfico**")
            rotulo_fonte = st.selectbox(
                "Fonte",
                options=list(FONTES_EXIBICAO),
                key="rotulo_fonte",
                format_func=lambda opcao: opcao,
                on_change=marcar_personalizado,
                help="Escolha entre estilos modernos, clássicos e monoespaçados, com variações de peso e inclinação.",
            )
            fonte = FONTES_EXIBICAO[rotulo_fonte]
            st.caption("As fontes são incorporadas como recursos padrão do PDF; não é preciso instalar nada.")
            tamanho_fonte = st.slider(
                "Tamanho da fonte",
                min_value=12,
                max_value=48,
                key="personalizacao_tamanho",
                help="Nomes compridos são reduzidos automaticamente para caber.",
                on_change=marcar_personalizado,
            )
            cor_texto = st.color_picker(
                "Cor do nome",
                key="personalizacao_cor",
                on_change=marcar_personalizado,
            )

    campos_extras = {}
    if campos_selecionados:
        with st.expander("Posição dos dados adicionais", expanded=False):
            st.caption("Cada campo será centralizado na posição escolhida, usando a fonte e a cor do nome.")
            for indice, campo in enumerate(campos_selecionados):
                st.markdown(f"**{campo.replace('_', ' ').capitalize()}**")
                chave_x = f"campo_extra_{campo}_x"
                chave_y = f"campo_extra_{campo}_y"
                st.session_state.setdefault(chave_x, 50)
                st.session_state.setdefault(chave_y, min(90, 62 + indice * 8))
                col_x, col_y = st.columns(2)
                with col_x:
                    st.slider("Posição horizontal (%)", 5, 95, key=chave_x, on_change=marcar_personalizado)
                with col_y:
                    st.slider("Posição vertical (%)", 5, 95, key=chave_y, on_change=marcar_personalizado)
                campos_extras[campo] = {
                    "x": st.session_state[chave_x],
                    "y": st.session_state[chave_y],
                    "fonte": fonte,
                    "tamanho": max(10, tamanho_fonte - 5),
                    "cor": cor_texto,
                }

    if arquivo_pdf:
        try:
            imagem_editor = imagem_primeira_pagina(arquivo_pdf.getvalue())
            areas_editor = {"nome": (posicao_x, posicao_y)}
            areas_editor.update({campo: (ajuste["x"], ajuste["y"]) for campo, ajuste in campos_extras.items()})
            resumo_pdf, avisos_pdf = diagnosticar_area_pdf(arquivo_pdf.getvalue(), areas_editor)
            st.caption(f"PDF validado · {resumo_pdf}")
            for aviso_pdf in avisos_pdf:
                st.warning(aviso_pdf)
            valor_editor = editor_posicao(
                imagem=imagem_editor,
                nome=amostra[0][coluna_nome] if amostra else "Nome do participante",
                x=posicao_x,
                y=posicao_y,
                tamanho=tamanho_fonte,
                cor=cor_texto,
                fonte=fonte,
                chave="editor-posicao-certificado",
                on_change=marcar_personalizado,
            )
            if valor_editor and valor_editor.get("nonce") != st.session_state.get("editor_posicao_nonce"):
                st.session_state["editor_posicao_nonce"] = valor_editor.get("nonce")
                st.session_state["personalizacao_x"] = int(valor_editor["x"])
                st.session_state["personalizacao_y"] = int(valor_editor["y"])
                posicao_x = st.session_state["personalizacao_x"]
                posicao_y = st.session_state["personalizacao_y"]
        except Exception as error:
            st.error(f"Não foi possível abrir o PDF no editor: {error}")
else:
    st.info("No PowerPoint, os nomes e demais dados seguem a posicao dos marcadores no proprio modelo.")

st.markdown('<div class="section-title">Marca visual e QR de verificação</div>', unsafe_allow_html=True)
logo_coluna, assinatura_coluna = st.columns(2)
with logo_coluna:
    arquivo_logo = st.file_uploader(
        "Logo opcional (PNG/JPG)", type=["png", "jpg", "jpeg"], key="logo_upload", max_upload_size=MAX_IMAGEM_MB
    )
with assinatura_coluna:
    arquivo_assinatura = st.file_uploader(
        "Assinatura opcional (PNG/JPG)", type=["png", "jpg", "jpeg"], key="assinatura_upload", max_upload_size=MAX_IMAGEM_MB
    )

imagens_extras = []
if arquivo_logo or arquivo_assinatura:
    with st.expander("Posição e tamanho das imagens", expanded=False):
        for rotulo, arquivo, x_padrao, y_padrao, largura_padrao in (
            ("Logo", arquivo_logo, 5, 5, 18),
            ("Assinatura", arquivo_assinatura, 70, 76, 20),
        ):
            if not arquivo:
                continue
            st.markdown(f"**{rotulo}**")
            prefixo = rotulo.lower()
            for sufixo, padrao in (("x", x_padrao), ("y", y_padrao), ("largura", largura_padrao)):
                st.session_state.setdefault(f"imagem_{prefixo}_{sufixo}", padrao)
            col_ix, col_iy, col_iw = st.columns(3)
            with col_ix:
                st.slider("Esquerda (%)", 0, 90, key=f"imagem_{prefixo}_x")
            with col_iy:
                st.slider("Topo (%)", 0, 90, key=f"imagem_{prefixo}_y")
            with col_iw:
                st.slider("Largura (%)", 5, 40, key=f"imagem_{prefixo}_largura")
            try:
                imagens_extras.append(
                    preparar_imagem(
                        arquivo,
                        st.session_state[f"imagem_{prefixo}_x"],
                        st.session_state[f"imagem_{prefixo}_y"],
                        st.session_state[f"imagem_{prefixo}_largura"],
                    )
                )
            except Exception as error:
                st.error(f"A imagem {rotulo.lower()} não pôde ser lida: {error}")

st.session_state.setdefault("incluir_qr", len(SEGREDO_QR) >= 32)
incluir_qr = st.checkbox(
    "Adicionar QR Code de verificação em cada certificado",
    key="incluir_qr",
    disabled=len(SEGREDO_QR) < 32,
    help="O QR inclui o nome, a data de emissao e os campos usados no modelo ou selecionados, assinados com HMAC.",
)
qr_habilitado = bool(incluir_qr and len(SEGREDO_QR) >= 32)
if len(SEGREDO_QR) < 32:
    st.warning("QR de verificação desativado. Configure `qr_secret` com pelo menos 32 caracteres nos Secrets do Streamlit.")
qr_posicao = {"x": 80, "y": 78, "largura": 10}
if qr_habilitado:
    st.caption("O QR contém o nome e os campos extras marcados; evite incluí-lo se esses dados não puderem ser compartilhados.")
    qx, qy, qt = st.columns(3)
    with qx:
        qr_posicao["x"] = st.slider("QR · esquerda (%)", 0, 80, 80)
    with qy:
        qr_posicao["y"] = st.slider("QR · topo (%)", 0, 80, 78)
    with qt:
        qr_posicao["largura"] = st.slider("QR · tamanho (%)", 6, 20, 10)

st.caption(
    f"Limites desta versão: planilha até {MAX_XLSX_MB} MB, PDF até {MAX_PDF_MB} MB, "
    f"até {MAX_PARTICIPANTES} participantes e saída estimada até {MAX_SAIDA_ESTIMADA_BYTES // (1024 * 1024)} MB."
)

fingerprint = (
    formato_saida,
    coluna_nome,
    posicao_x,
    posicao_y,
    fonte,
    tamanho_fonte,
    cor_texto,
    tuple(
        (campo, tuple(sorted(ajuste.items())))
        for campo, ajuste in sorted(campos_extras.items())
    ),
    sha256(arquivo_planilha.getvalue()).hexdigest() if arquivo_planilha else None,
    sha256(arquivo_modelo.getvalue()).hexdigest() if arquivo_modelo else None,
    tuple(
        (sha256(imagem["dados"]).hexdigest(), imagem["x"], imagem["y"], imagem["largura"])
        for imagem in imagens_extras
    ),
    qr_habilitado,
    tuple(sorted(qr_posicao.items())),
)
if st.session_state.get("fingerprint") != fingerprint:
    st.session_state.pop("zip_certificados", None)
    st.session_state.pop("pdf_unico", None)
    st.session_state.pop("pdf_previa", None)
    st.session_state.pop("total_certificados", None)
    st.session_state.pop("formato_resultado", None)
    st.session_state["fingerprint"] = fingerprint

if arquivo_pdf and amostra and amostra[0].get(coluna_nome, "").strip():
    try:
        with st.spinner("Preparando prévia com o primeiro participante…"):
            if "pdf_previa" not in st.session_state:
                emissao_previa = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M")
                qr_previa = (
                    qr_participante_png(amostra[0], "PDF", emissao_previa, list(campos_selecionados))
                    if qr_habilitado
                    else None
                )
                st.session_state["pdf_previa"] = gerar_previa_pdf(
                    arquivo_pdf.getvalue(),
                    amostra[0],
                    posicao_x,
                    posicao_y,
                    fonte,
                    tamanho_fonte,
                    cor_texto,
                    campos_extras,
                    imagens_extras,
                    qr_previa,
                    qr_posicao,
                )
        with st.expander("Prévia do primeiro certificado", expanded=True):
            dados_pdf = base64.b64encode(st.session_state["pdf_previa"]).decode("ascii")
            st.iframe(
                f"data:application/pdf;base64,{dados_pdf}",
                height=540,
                alt=f"Prévia do certificado de {amostra[0][coluna_nome]}",
            )
    except Exception as error:
        st.warning(f"Não foi possível criar a prévia: {error}")

st.markdown("<div style='height:.5rem'></div>", unsafe_allow_html=True)
if st.button("Gerar certificados", type="primary", width="stretch"):
    if not arquivo_planilha or not arquivo_modelo:
        st.error("Envie a planilha e o modelo de certificado antes de continuar.")
    elif erro_dados or (formato_saida == "PowerPoint" and erro_modelo_pptx):
        st.error(erro_dados or erro_modelo_pptx)
    else:
        try:
            with tempfile.TemporaryDirectory(prefix="gerador_certificados_") as diretorio:
                pasta = Path(diretorio)
                caminho_planilha = pasta / "participantes.xlsx"
                caminho_modelo = pasta / ("certificado.pdf" if formato_saida == "PDF" else "modelo.pptx")
                caminho_planilha.write_bytes(arquivo_planilha.getvalue())
                caminho_modelo.write_bytes(arquivo_modelo.getvalue())
                participantes = ler_participantes(caminho_planilha)
                validar_participantes(participantes, coluna_nome, campos=campos_selecionados)
                if len(participantes) > MAX_PARTICIPANTES:
                    raise ValueError(f"Limite de {MAX_PARTICIPANTES} participantes por geracao; a planilha contem {len(participantes)}.")
                quantidade_slides = len(Presentation(caminho_modelo).slides) if formato_saida == "PowerPoint" else 1
                sobrecarga = (64 * 1024 if fonte == "GreatVibes" else 8 * 1024) if formato_saida == "PDF" else 0
                sobrecarga += sum(len(imagem["dados"]) for imagem in imagens_extras) * quantidade_slides
                if qr_habilitado:
                    sobrecarga += 24 * 1024 * quantidade_slides
                saida_estimada = len(participantes) * (arquivo_modelo.size + sobrecarga)
                if saida_estimada > MAX_SAIDA_ESTIMADA_BYTES:
                    limite_mb = MAX_SAIDA_ESTIMADA_BYTES // (1024 * 1024)
                    raise ValueError(f"A saida estimada passa de {limite_mb} MB. Use um modelo menor ou divida a lista em lotes.")
                emissao = datetime.now(ZoneInfo("America/Sao_Paulo")).strftime("%d/%m/%Y %H:%M")
                campos_qr = list(campos_selecionados)
                if formato_saida == "PowerPoint":
                    campos_qr = sorted(marcadores_modelo - {"nome"})

                def qr_para_participante(pessoa: dict[str, str]) -> bytes:
                    return qr_participante_png(pessoa, formato_saida, emissao, campos_qr)

                if formato_saida == "PDF":
                    progresso = st.progress(0, text="Preparando os certificados")
                    def atualizar_progresso(atual: int, total: int) -> None:
                        progresso.progress(int(atual / total * 100), text=f"Gerando certificado {atual} de {total}")
                    try:
                        arquivos = gerar_certificados_pdf(
                            caminho_modelo, pasta / "certificados", participantes, coluna_nome,
                            posicao_x=posicao_x, posicao_y=posicao_y, fonte=fonte,
                            tamanho_fonte=tamanho_fonte, cor=cor_texto,
                            ao_gerar=atualizar_progresso, campos_extras=campos_extras,
                            imagens=imagens_extras,
                            qr_por_participante=qr_para_participante if qr_habilitado else None,
                            qr_posicao=qr_posicao,
                        )
                    finally:
                        progresso.empty()
                    pdf_unico = combinar_certificados_pdf(arquivos)
                    extensao = "PDF"
                else:
                    arquivos = []
                    pasta_saida = pasta / "certificados"
                    pasta_saida.mkdir()
                    for indice, pessoa in enumerate(participantes, start=1):
                        imagens_pessoa = list(imagens_extras)
                        if qr_habilitado:
                            imagens_pessoa.append({"dados": qr_para_participante(pessoa), **qr_posicao})
                        destino = pasta_saida / f"{indice:03d}_{nome_seguro(pessoa[coluna_nome])}.pptx"
                        preencher_modelo(caminho_modelo, destino, pessoa, imagens=imagens_pessoa)
                        arquivos.append(destino)
                    pdf_unico = None
                    extensao = "PowerPoint"

                arquivo_zip = criar_zip(arquivos)
                st.session_state["zip_certificados"] = arquivo_zip
                st.session_state["pdf_unico"] = pdf_unico
                st.session_state["total_certificados"] = len(arquivos)
                st.session_state["formato_resultado"] = extensao
                guardar_no_historico(extensao, len(arquivos), arquivo_zip)
        except (FileNotFoundError, ValueError, PermissionError, OSError) as error:
            st.error(f"Nao foi possivel gerar os certificados: {error}")
        except Exception as error:
            st.error(f"Ocorreu um erro ao processar os arquivos: {error}")

if st.session_state.get("zip_certificados"):
    formato_resultado = st.session_state.get("formato_resultado", "PDF")
    st.success(f"{st.session_state['total_certificados']} certificados {formato_resultado} prontos.")
    st.download_button(
        f"Baixar arquivos {formato_resultado} (.zip)",
        data=st.session_state["zip_certificados"],
        file_name="certificados.zip" if formato_resultado == "PDF" else "certificados_powerpoint.zip",
        mime="application/zip",
        type="primary",
        width="stretch",
    )
    if st.session_state.get("pdf_unico"):
        st.download_button(
            "Baixar todos em um PDF",
            data=st.session_state["pdf_unico"],
            file_name="certificados_completos.pdf",
            mime="application/pdf",
            width="stretch",
        )

if st.session_state.get("historico_emissoes"):
    st.subheader("Historico desta sessao")
    st.caption("Os lotes ficam disponiveis enquanto esta sessao do navegador estiver aberta.")
    for lote in st.session_state["historico_emissoes"]:
        with st.expander(f"{lote['data']} - {lote['quantidade']} certificados - {lote['formato']}"):
            if lote.get("zip"):
                st.download_button(
                    "Baixar novamente este lote (.zip)",
                    data=lote["zip"],
                    file_name=f"certificados_{lote['id']}.zip",
                    mime="application/zip",
                    key=f"historico_{lote['id']}",
                    width="stretch",
                )
            else:
                st.info("Este lote ultrapassou o limite reservado ao historico; gere-o novamente para baixar.")
    if st.button("Limpar historico desta sessao", key="limpar_historico"):
        st.session_state["historico_emissoes"] = []
        st.rerun()

st.markdown(
    '<div class="privacy-note">Seus arquivos são processados para esta geração e não devem conter dados confidenciais.</div>',
    unsafe_allow_html=True,
)
